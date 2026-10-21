import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r'C:\Users\86136\Desktop\cv_exp2'
OUT = os.path.join(BASE, '实验二_图像增强_实验报告.docx')

doc = docx.Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(3.17)
sec.right_margin = Cm(3.17)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def set_run(run, size=12, bold=False, name='Times New Roman', east='宋体'):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    return run


def para(text='', indent=True, align=None, before=0, after=0, line=1.5, bold=False, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    if text:
        set_run(p.add_run(text), size, bold)
    return p


def heading(text, level=1):
    if level == 1:
        return para(text, indent=False, before=10, after=4, line=1.3, bold=True, size=14)
    return para(text, indent=False, before=6, after=2, line=1.3, bold=True, size=12)


def code(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.left_indent = Cm(0.6)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F4F4F4')
    p._p.get_or_add_pPr().append(shd)
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        set_run(r, 9, False, 'Consolas', '宋体')
        if i < len(lines) - 1:
            r.add_break()
    return p


def figure(path, width, caption):
    doc.add_picture(path, width=Cm(width))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    cap = para(caption, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
    for r in cap.runs:
        set_run(r, 10.5)


para('实验二：图像增强', indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, line=1.2, bold=True, size=16)
para('姓名：李铱璠 学号：202410315057', indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)

heading('一、实验目的')
para('熟悉OpenCV的基本用法，学会对图像做平滑和滤波，完成简单的图像增强。同时自己动手写中值滤波和高斯卷积，'
     '弄清楚滤波器内部在算什么，以及不同噪声该用哪种滤波处理。')

heading('二、实验内容')

heading('2.1 导入依赖库', 2)
code('''import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import cv2
from numpy.lib.stride_tricks import sliding_window_view''')
para('cv2用来读写图像和调用滤波函数，numpy做数组运算，matplotlib把结果画出来。脚本要在没有图形界面的环境里跑，'
     '所以用Agg后端，图片直接savefig保存，不调用plt.show。')

heading('2.2 读取原始图像与色彩空间转换', 2)
code('''img = cv2.imread('p1.jpg')
print('图像尺寸:', img.shape)

(b, g, r) = img[100, 100]
print('像素(100,100) BGR =', b, g, r)

rgb_img  = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
gray_img = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)''')
para('运行结果是图像尺寸(1104,1104,3)，(100,100)处的BGR值为(73,84,242)。R分量是242，说明这个点落在左上角的红色背景上。')
para('OpenCV读进来的通道顺序是BGR，matplotlib却按RGB解释，直接imshow会把红蓝两个通道对调，画面变成蓝底、衣服发黄，'
     '也就是图1左边那张。用cvtColor转成RGB以后颜色恢复正常，见图1中间。转灰度是把三个通道按亮度加权合成一个通道，'
     '只剩明暗信息，后面看噪声和滤波效果更方便，见图1右边。')
figure(os.path.join(BASE, 'output_images', 'color_space.jpg'), 15.2,
       '图1 原始图像（BGR直接显示）、RGB图像与灰度图像')

heading('2.3 添加噪声', 2)
code('''def add_salt_pepper(image, amount=0.4):
    out = image.copy()
    h, w = out.shape[:2]
    n = int(amount * h * w)
    ys = np.random.randint(0, h, n)
    xs = np.random.randint(0, w, n)
    half = n // 2
    out[ys[:half], xs[:half]] = 0.0
    out[ys[half:], xs[half:]] = 1.0
    return out


def add_gaussian_noise(image, mean=0.2, var=0.03):
    noise = np.random.normal(mean, np.sqrt(var), image.shape)
    return np.clip(image + noise, 0.0, 1.0)''')
para('椒盐噪声从图里随机挑40%的像素，一半改成纯黑，一半改成纯白，看上去是一些散落的黑白点。'
     '高斯噪声是给每个像素加一个服从N(0.2,0.03)的随机数，整幅图都有颗粒感；均值取0.2不是0，画面会整体变亮。'
     '两种噪声不太一样，椒盐只动被选中的少数像素，高斯噪声每个像素都受影响。')
figure(os.path.join(BASE, 'output_images', 'noise_comparison.jpg'), 15.2,
       '图2 原图、椒盐噪声与高斯噪声对比')

heading('2.4 三种滤波方法对比', 2)
code('''mean_sp  = cv2.blur(sp_noise_img, (5, 5))
mid_sp   = cv2.medianBlur((sp_noise_img * 255).astype(np.uint8), 5)
gauss_sp = cv2.GaussianBlur((sp_noise_img * 255).astype(np.uint8), (5, 5), 0)
# 高斯噪声同理，得到mean_gus / mid_gus / gauss_gus''')
para('均值滤波把邻域内所有像素求平均，再替换中心像素。椒盐噪声的黑白点会被摊成灰斑，去不干净，边缘也会一起变模糊，'
     '见图3左上。中值滤波取邻域像素的中位数，只要黑白点不到邻域的一半就会被完全排掉，所以椒盐噪声基本能恢复干净，'
     '边缘也保得住，见图3中上。高斯滤波按距离加权平均，适合处理连续的高斯噪声，平滑得比较自然，'
     '但椒盐点只会被糊成小团，不会彻底消失，见图3右上。')
para('第二行是对高斯噪声做同样三种处理。三种方法都能压住噪声，均值和高斯滤波更平滑，中值滤波留下的颗粒稍多，'
     '不过边缘细节也更多，见图3下排。')
figure(os.path.join(BASE, 'result', 'filter_results_2x3.jpg'), 14.6,
       '图3 两种噪声下三种滤波方法的效果对比')

heading('2.5 手动实现彩色中值滤波', 2)
code('''def manual_median_filter_color(image, kernel_size=5):
    pad = kernel_size // 2
    out = np.zeros_like(image, dtype=np.uint8)
    padded = np.pad(image, ((pad, pad), (pad, pad), (0, 0)), mode='edge')
    h, w = image.shape[:2]
    for r0 in range(0, h, 64):
        r1 = min(r0 + 64, h)
        win = sliding_window_view(padded[r0:r1 + 2 * pad],
                                  (kernel_size, kernel_size), axis=(0, 1))
        out[r0:r1] = np.median(win, axis=(-2, -1)).astype(np.uint8)
    return out''')
para('写的时候注意两点。一是先对图像做edge填充，不然边界像素凑不齐5×5的邻域。二是按R、G、B三个通道分别取中值，'
     '不能对三维数组整体取，否则三个通道的数值混在一起，颜色会乱掉。直接三重循环太慢，这里用sliding_window_view'
     '把邻域取出来，再按64行分块算，整个过程0.76秒。和cv2.medianBlur逐像素比较，最大差值是0，'
     '说明手写的结果和OpenCV完全一样。')
figure(os.path.join(BASE, 'result', 'manual_median.jpg'), 14.6,
       '图4 椒盐噪声图像与手动中值滤波结果（k=5）')

heading('2.6 自定义高斯核与二维卷积', 2)
code('''def gauss(kernel_size, sigma):
    kernel = np.zeros((kernel_size, kernel_size), dtype=np.float64)
    center = kernel_size // 2
    s = sigma ** 2
    for i in range(kernel_size):
        for j in range(kernel_size):
            x, y = i - center, j - center
            kernel[i, j] = np.exp(-(x ** 2 + y ** 2) / (2 * s))
    return kernel / kernel.sum()

k_gauss = gauss(7, 2)
conv_gus = cv2.filter2D(gus_noise_img.astype(np.float32), -1, k_gauss.astype(np.float32))''')
para('按二维高斯函数把7×7的权重一个个算出来，再归一化，让核里所有元素加起来等于1。不归一化的话卷积后整幅图会变亮或变暗。'
     '这个核中心权重约0.0467，越往外越小。用filter2D卷完，高斯噪声明显被抹平了。核越大、sigma越大，图就越模糊。')
figure(os.path.join(BASE, 'result', 'custom_gaussian.jpg'), 15.0,
       '图5 高斯噪声、自定义7×7高斯核与卷积结果')

heading('三、实验结果与分析')
para('（1）同一张图在BGR、RGB、灰度三种表示下看起来完全不同，说明通道顺序会影响显示结果。实际使用时读进来先转成RGB再显示，'
     '可以避免红蓝互换。')
para('（2）椒盐噪声只影响被选中的像素，是稀疏的极端值；高斯噪声作用在所有像素上，是稠密的连续扰动。'
     '噪声类型不同，后面选的滤波方法也不同。')
para('（3）椒盐噪声用中值滤波最合适，高斯噪声用均值或高斯滤波更合适。实验结果和这个判断一致，中值滤波对椒盐噪声恢复得最干净，'
     '均值和高斯滤波对高斯噪声平滑得最自然。')
para('（4）手写的中值滤波和OpenCV输出逐像素相同，最大差值0，说明中值滤波本质上就是邻域排序后取中间那个值。'
     '自定义高斯核也能复现GaussianBlur的效果，说明高斯滤波就是按距离加权的邻域平均。')
para('（5）核从3×3加到5×5、7×7，去噪更彻底，但边缘也更模糊。椒盐噪声比例设成0.4已经比较重，中值滤波还能恢复，'
     '说明它对极端值不敏感。')

heading('四、实验小结')
para('1.遇到的问题及解决方法', indent=False, bold=True)
para('（1）cv2.imread读中文路径会返回None。把图片和脚本放到纯英文目录，再用os.path.dirname(os.path.abspath(__file__))'
     '拼绝对路径，换目录也能跑。')
para('（2）环境里没有scikit-image，random_noise导不进来。就用numpy自己写椒盐和高斯噪声，顺便把两种噪声的公式弄清楚了。')
para('（3）对(H,W,3)整体取中值颜色会错。改成三个通道分开滤波再合起来，结果才和OpenCV对得上。')
para('（4）一开始用三重循环，跑得太慢。换成滑动窗口加按行分块，降到0.76秒，结果没变。')
para('2.收获', indent=False, bold=True)
para('这次实验把OpenCV的读取、色彩转换、加噪声和三种滤波都过了一遍，也自己写了中值滤波和高斯卷积。'
     '写完以后对“滤波就是邻域里做加权或者排序”这句话有了具体的感觉。另外去噪方法没有绝对的好坏，'
     '关键看噪声是什么类型，这一点对后面的边缘检测和形态学处理也有用。')

doc.save(OUT)
print('SAVED', OUT)
